import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from docx import Document

from src.generate_report import build_report


class ReportContentTests(unittest.TestCase):
    def test_report_stays_within_evidence_boundary(self):
        with TemporaryDirectory() as directory, patch("src.generate_report.OUTPUT", Path(directory)):
            document = Document(build_report())
        text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )

        for objective in (
            "Identify Associated Factors",
            "Compare Available Stressor Evidence",
            "Prioritise Field Verification",
            "Frame Conservation with Economic Context",
        ):
            self.assertIn(objective, text)

        self.assertIn("RM8.7 billion", text)
        self.assertIn("six evaluated archipelagos", text)
        self.assertIn("2011-2015", text)
        self.assertIn("NOAA variables are excluded from the scored model", text)
        self.assertIn("2025", text)

        for unsupported in (
            "RM 47.73",
            "78%",
            "RM 842.5",
            "20-year NPV",
            "Mandate Diver Quotas",
            "carrying-capacity quota",
        ):
            self.assertNotIn(unsupported, text)


if __name__ == "__main__":
    unittest.main()
