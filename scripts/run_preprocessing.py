import os
import json
import sys
import polars as pl
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.pipeline import parse_noaa_row

base_dir = str(ROOT)
raw_dir = os.path.join(base_dir, "data", "raw")
proc_dir = os.path.join(base_dir, "data", "processed")
os.makedirs(proc_dir, exist_ok=True)

print("================================================================")
print("STEP 1: INSPECT RAW STRUCTURED DATA BEFORE COMBINING")
print("================================================================")

import openpyxl

rc_path = os.path.join(raw_dir, "structured", "reef_check", "ReefCheck_Malaysia_FINAL.xlsx")
wb = openpyxl.load_workbook(rc_path, data_only=True)
ws = wb["model_ready"]
data = list(ws.iter_rows(values_only=True))
headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(data[0])]
rows = data[1:]
cols = {h: [r[i] for r in rows] for i, h in enumerate(headers)}
df_rc = pl.DataFrame(cols, strict=False)
df_rc = df_rc.with_columns(
    pl.when(pl.col("island") == "Labuan")
    .then(pl.lit("W.P. Labuan"))
    .otherwise(pl.col("state"))
    .alias("state")
)
print(f"\n[1] Reef Check Survey Panel (model_ready): {df_rc.shape[0]} rows, {df_rc.shape[1]} columns")
print("Columns:", df_rc.columns[:10], "... (total", len(df_rc.columns), ")")
print("Null counts per column BEFORE merge:")
for col in df_rc.columns:
    nulls = df_rc[col].null_count()
    if nulls > 0:
        print(f"  - {col:28s}: {nulls:3d} nulls ({nulls/len(df_rc)*100:.1f}%)")
    else:
        print(f"  - {col:28s}:   0 nulls (Clean)")

# 2. JUPEM Geocoding Reference
geo_path = os.path.join(raw_dir, "structured", "geocoding", "island_coordinates.csv")
df_geo = pl.read_csv(geo_path, encoding="utf8")
print(f"\n[2] Island Coordinates (JUPEM): {df_geo.shape[0]} rows, {df_geo.shape[1]} columns")
print("Columns:", df_geo.columns)
print("Null counts in Coordinates:")
for col in df_geo.columns:
    print(f"  - {col:20s}: {df_geo[col].null_count()} nulls")

# 2b. Island Built-Environment & Accommodation Capacity
infr_path = os.path.join(raw_dir, "structured", "infrastructure", "island_accommodations.csv")
df_infr = pl.read_csv(infr_path, encoding="utf8")
print(f"\n[2b] Island Accommodation Capacity: {df_infr.shape[0]} rows, {df_infr.shape[1]} columns")
print("Null counts in Accommodations:")
for col in df_infr.columns:
    print(f"  - {col:25s}: {df_infr[col].null_count()} nulls")

# 3. NOAA Coral Reef Watch Stations
print(f"\n[3] NOAA Coral Reef Watch Virtual Stations:")
noaa_dir = os.path.join(raw_dir, "structured", "noaa_crw")
station_files = {
    "malacca_strait": "malacca_strait.txt",
    "sabah": "sabah.txt",
    "northern_borneo": "northern_borneo.txt",
    "singapore": "singapore.txt",
    "west_gulf_of_thailand": "west_gulf_of_thailand.txt"
}

# Assign station mapping based on state/ecoregion
def get_station_id(state, island):
    isl_lower = island.lower()
    st_lower = state.lower() if state else ""
    if "sabah" in st_lower or any(name in isl_lower for name in ("sipadan", "mabul", "kudat", "labuan")):
        return "sabah"
    elif "sarawak" in st_lower:
        return "northern_borneo"
    elif "kedah" in st_lower or "perak" in st_lower or "payar" in isl_lower or "pangkor" in isl_lower:
        return "malacca_strait"
    elif "johor" in st_lower or "tioman" in isl_lower or "aur" in isl_lower or "sibu" in isl_lower or "tinggi" in isl_lower:
        return "singapore"
    else:
        return "west_gulf_of_thailand"

for st_id, fname in station_files.items():
    fpath = os.path.join(noaa_dir, fname)
    sz_kb = os.path.getsize(fpath) / 1024
    print(f"  - Station '{st_id}': {fname} ({sz_kb:.1f} KB)")

# Parse NOAA annual thermal metrics
noaa_annual = {}
for st_id, fname in station_files.items():
    fpath = os.path.join(noaa_dir, fname)
    with open(fpath, "r") as f:
        lines = f.readlines()
    data_lines = [l.strip().split() for l in lines if l.strip() and not l.startswith("#") and not l.startswith("YYYY")]
    by_year = {}
    for parts in data_lines:
        try:
            yr, ssta, dhw = parse_noaa_row(parts)
            if yr not in by_year:
                by_year[yr] = {"max_dhw": dhw, "ssta_list": [ssta]}
            else:
                by_year[yr]["max_dhw"] = max(by_year[yr]["max_dhw"], dhw)
                by_year[yr]["ssta_list"].append(ssta)
        except (ValueError, IndexError):
            continue
    for yr, vals in by_year.items():
        mean_ssta = sum(vals["ssta_list"]) / len(vals["ssta_list"])
        noaa_annual[(st_id, yr)] = {
            "noaa_max_dhw": round(vals["max_dhw"], 2),
            "noaa_mean_ssta": round(mean_ssta, 2)
        }

print(f"  -> Extracted {len(noaa_annual)} station-year annual thermal stress profiles.")

# 4. OpenDOSM Datasets
print(f"\n[4] OpenDOSM Datasets:")
gdp_path = os.path.join(raw_dir, "structured", "opendosm", "gdp_state_real_supply.json")
with open(gdp_path, "r") as f:
    gdp_json = json.load(f)
print(f"  - GDP State Real Supply: {len(gdp_json)} records")

fish_path = os.path.join(raw_dir, "structured", "opendosm", "fish_landings.csv")
df_fish = pl.read_csv(fish_path)
print(f"  - Marine Fish Landings: {df_fish.shape[0]} records across states {df_fish['state'].n_unique()}")

water_path = os.path.join(raw_dir, "structured", "opendosm", "water_pollution_basin.csv")
df_water = pl.read_csv(water_path)
print(f"  - River Basin Water Quality: {df_water.shape[0]} records")

print("\n================================================================")
print("STEP 2: STEP-BY-STEP DATA COMBINATION & HARMONIZATION")
print("================================================================")

# Join 1: Reef Check + Geocoding Coordinates
# Check island column name and overlap
geo_clean = df_geo.select(["island", "latitude", "longitude", "marine_park"])
df_merged = df_rc.join(geo_clean, on="island", how="left")
print(f"Join 1 (RC + JUPEM Geocoding): {df_merged.shape[0]} rows, {df_merged.shape[1]} columns")
missing_coords = df_merged["latitude"].null_count()
print(f"  - Missing Coordinates count: {missing_coords} ({'CLEAN 100%' if missing_coords==0 else 'WARNING'})")

# Join 2: Merge NOAA Thermal Stress
max_dhw_list = []
mean_ssta_list = []
station_col = []
for row in df_merged.iter_rows(named=True):
    isl = str(row.get("island", ""))
    st = str(row.get("state", ""))
    yr = int(row.get("survey_year", 2020))
    st_id = get_station_id(st, isl)
    station_col.append(st_id)
    metrics = noaa_annual.get((st_id, yr), {"noaa_max_dhw": None, "noaa_mean_ssta": None})
    max_dhw_list.append(metrics["noaa_max_dhw"])
    mean_ssta_list.append(metrics["noaa_mean_ssta"])

df_merged = df_merged.with_columns([
    pl.Series("noaa_station_id", station_col),
    pl.Series("noaa_max_dhw", max_dhw_list),
    pl.Series("noaa_mean_ssta", mean_ssta_list)
])
print(f"Join 2 (+ NOAA Thermal Stress): {df_merged.shape[0]} rows, {df_merged.shape[1]} columns")

# Join 3: Merge Island Built-Environment & Accommodation Capacity
infr_clean = df_infr.select([
    "island", 
    "resort_count", 
    "estimated_room_capacity", 
    "dive_center_count", 
    "has_commercial_jetty"
])
df_merged = df_merged.join(infr_clean, on="island", how="left").with_columns([
    pl.col("resort_count").fill_null(0),
    pl.col("estimated_room_capacity").fill_null(0),
    pl.col("dive_center_count").fill_null(0),
    pl.col("has_commercial_jetty").fill_null(0)
])
print(f"Join 3 (+ Accommodation Capacity): {df_merged.shape[0]} rows, {df_merged.shape[1]} columns")

print("OpenDOSM context files retained for descriptive use; no island-level tourism values were inferred.")

print("\n================================================================")
print("STEP 3: POST-COMBINATION AUDIT & DATA VERIFICATION")
print("================================================================")
print(f"Final Merged Dataset Dimensions: {df_merged.shape[0]} rows x {df_merged.shape[1]} columns")
print("\nPost-Combination Null Audit:")
for col in df_merged.columns:
    nulls = df_merged[col].null_count()
    status = "OK (0 nulls)" if nulls == 0 else f"{nulls} nulls ({nulls/len(df_merged)*100:.1f}%)"
    print(f"  - {col:30s}: {status}")

# Save master dataset
csv_dest = os.path.join(proc_dir, "master_reef_tourism_dataset.csv")
df_merged.write_csv(csv_dest)
print(f"\nSuccessfully exported: {csv_dest} ({os.path.getsize(csv_dest)/1024:.1f} KB)")
